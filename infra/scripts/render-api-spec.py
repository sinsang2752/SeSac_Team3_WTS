#!/usr/bin/env python3
"""
openapi.yaml 을 사람이 읽는 표 형식 API 명세서로 바꾼다.

    python3 infra/scripts/render-api-spec.py docs/api/openapi.yaml

산출물 (yaml 과 같은 폴더):
    api-spec.md     GitHub에서 바로 읽는다
    api-spec.xlsx   제출·공유용
    api-spec.pdf    인쇄·배포용. reportlab 과 한글 TrueType 글꼴이 있을 때만 만든다

모두 생성물이다. 직접 고치지 않는다. 원본은 컨트롤러(@Operation)와 DTO(@Schema)이고,
보통은 generate-openapi.sh 가 yaml 을 만든 뒤 이 스크립트를 부른다.
yaml 만 읽으므로 서비스를 띄우지 않고도 돌릴 수 있다.
"""
import json
import math
import os
import sys

try:
    import yaml
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.hyperlink import Hyperlink
except ImportError as e:
    sys.exit(f"{e.name} 모듈이 필요하다: pip install pyyaml openpyxl")

# 목차 순서. 사용 흐름(로그인 → 종목·시세 → 계좌·주문 → 결과)을 따른다. 여기 없는 태그는 뒤에 붙는다.
TAG_ORDER = ["사용자", "관심종목", "종목", "시세", "차트", "계좌", "주문", "체결", "보유종목", "포트폴리오"]
METHOD_ORDER = {"post": 0, "get": 1, "put": 2, "patch": 3, "delete": 4}
STATUS_TEXT = {200: "OK", 201: "Created", 204: "No Content"}

AUTH_ROW = {"name": "Authorization", "in": "header", "type": "string", "required": True,
            "description": "Bearer {accessToken} — mock-login 이 발급한 토큰", "example": ""}


# ── 모델 ──────────────────────────────────────────────────────────

class Spec:
    def __init__(self, doc):
        self.doc = doc
        self.schemas = doc["components"]["schemas"]
        self.error_codes = {e["code"]: e for e in doc.get("x-error-codes", [])}
        self.base_url = doc["servers"][0]["url"]
        self.version = doc["info"]["version"]
        self.ops = self._operations()

    def resolve(self, s):
        while "$ref" in s:
            s = self.schemas[s["$ref"].rsplit("/", 1)[-1]]
        return s

    @staticmethod
    def ref_name(s):
        if "$ref" in s:
            return s["$ref"].rsplit("/", 1)[-1]
        if s.get("type") == "array" and "$ref" in s.get("items", {}):
            return s["items"]["$ref"].rsplit("/", 1)[-1]
        return None

    def type_label(self, s):
        if "$ref" in s:
            return "object"
        if "enum" in s:
            return "enum (" + " / ".join(map(str, s["enum"])) + ")"
        t = s.get("type")
        if isinstance(t, list):  # OpenAPI 3.1: ["string", "null"]
            t = next((x for x in t if x != "null"), t[0])
        if t == "array":
            return f"array<{self.type_label(s.get('items', {}))}>"
        if t == "string" and s.get("format") == "date-time":
            return "string (date-time)"
        return t or "object"

    def fields(self, schema, prefix=""):
        """스키마를 평탄화한다. 중첩 객체는 a.b, 배열 원소는 a[].b 로 적는다."""
        s = self.resolve(schema)
        if s.get("type") == "array":
            return self.fields(s.get("items", {}), prefix)
        rows = []
        required = set(s.get("required", []))
        for name, prop in (s.get("properties") or {}).items():
            rows.append({
                "name": prefix + name,
                "type": self.type_label(prop),
                "required": name in required,
                "description": prop.get("description", ""),
                "example": prop.get("example", ""),
            })
            inner = self.resolve(prop.get("items", {}) if prop.get("type") == "array" else prop)
            if inner.get("properties"):
                child = prefix + name + ("[]." if prop.get("type") == "array" else ".")
                rows.extend(self.fields(inner, child))
        return rows

    def example(self, schema):
        """예시 JSON. 예시가 없는 단순 필드는 뺀다(조건부 필드가 모순된 예시를 만들지 않게)."""
        s = self.resolve(schema)
        if s.get("type") == "array":
            return [self.example(s.get("items", {}))]
        out = {}
        for name, prop in (s.get("properties") or {}).items():
            if "example" in prop:
                out[name] = prop["example"]
            elif prop.get("type") == "array" or "$ref" in prop or prop.get("properties"):
                out[name] = self.example(prop)
        return out

    def _operations(self):
        tag_rank = {t: i for i, t in enumerate(TAG_ORDER)}
        for i, t in enumerate(x["name"] for x in self.doc.get("tags", [])):
            tag_rank.setdefault(t, len(TAG_ORDER) + i)
        tag_desc = {t["name"]: t.get("description", "") for t in self.doc.get("tags", [])}

        ops = []
        for path, items in self.doc["paths"].items():
            for method, op in items.items():
                if method not in METHOD_ORDER:
                    continue
                tag = (op.get("tags") or ["기타"])[0]
                success = next(code for code in sorted(op["responses"]) if code.startswith("2"))
                body_schema = (op["responses"][success].get("content") or {}) \
                    .get("application/json", {}).get("schema")
                req_schema = (op.get("requestBody", {}).get("content") or {}) \
                    .get("application/json", {}).get("schema")
                ops.append({
                    "tag": tag, "tag_description": tag_desc.get(tag, ""),
                    "summary": op["summary"], "description": op.get("description", ""),
                    "method": method.upper(), "path": path,
                    "public": op.get("security") == [],
                    "params": op.get("parameters", []),
                    "request_schema": req_schema,
                    "status": int(success), "response_schema": body_schema,
                    "errors": op.get("x-error-codes", []),
                    "_sort": (tag_rank.setdefault(tag, 999), path.count("/"),
                              METHOD_ORDER[method], path),
                })
        ops.sort(key=lambda o: o["_sort"])
        for no, op in enumerate(ops, 1):
            op["no"] = no
        return ops

    def request_rows(self, op):
        rows = [] if op["public"] else [AUTH_ROW]
        for p in op["params"]:
            rows.append({"name": p["name"], "in": p["in"], "type": self.type_label(p.get("schema", {})),
                         "required": p.get("required", False),
                         "description": p.get("description", "")
                         + (f" (기본값 {p['schema']['default']})" if "default" in p.get("schema", {}) else ""),
                         "example": p.get("example", "")})
        if op["request_schema"]:
            rows += [dict(r, **{"in": "body"}) for r in self.fields(op["request_schema"])]
        return rows

    def response_label(self, op):
        status = f"{op['status']} {STATUS_TEXT.get(op['status'], '')}".strip()
        s = op["response_schema"]
        if not s:
            return status + " · 본문 없음"
        name = self.ref_name(s)
        return status + " · " + (f"{name} 배열" if s.get("type") == "array" else name)

    def errors_of(self, op):
        rows = [self.error_codes[c] for c in op["errors"]]
        return sorted(rows, key=lambda e: e["status"])  # 안정 정렬이라 같은 상태코드 안의 순서는 유지된다


def cell_text(v):
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    return str(v)


def pretty(obj):
    return json.dumps(obj, ensure_ascii=False, indent=2)


# ── Markdown ──────────────────────────────────────────────────────

def esc(v):
    """표 칸의 일반 텍스트. array<object> 같은 꺾쇠가 HTML 태그로 해석되지 않게 한다."""
    t = cell_text(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return t.replace("|", "\\|").replace("\n", "<br>")


def code(v):
    t = cell_text(v)
    return "`" + t.replace("|", "\\|") + "`" if t else ""


def md_table(headers, rows):
    """rows 의 칸은 esc() 또는 code() 로 이미 만든 문자열이다."""
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(out)


def render_md(spec):
    L = []
    w = L.append
    w("# WTS 모의투자 API 명세서\n")
    w("> **생성물이다. 직접 고치지 않는다.** 원본은 컨트롤러(`@Operation`)와 DTO(`@Schema`)이고, "
      "[`openapi.yaml`](openapi.yaml)을 거쳐 `./infra/scripts/generate-openapi.sh` 가 만든다. "
      "같은 내용의 엑셀·PDF는 [`api-spec.xlsx`](api-spec.xlsx), [`api-spec.pdf`](api-spec.pdf).\n>\n"
      "> 실시간 시세 WebSocket(`/ws/market`)과 Kafka 이벤트 계약은 OpenAPI로 표현할 수 없어 "
      "[README.md](README.md)에 있다.\n")
    w(f"버전 `{spec.version}` · API {len(spec.ops)}개\n")

    w("## 1. 공통\n")
    # 아래 두 표는 손으로 쓴 마크다운 고정 문구라 esc() 를 거치지 않는다.
    w(md_table(["항목", "내용"], [
        ["Base URL", f"`{spec.base_url}` (Gateway). 각 서비스를 직접 호출하지 않는다"],
        ["인증", "`Authorization: Bearer <accessToken>`. 토큰은 `POST /api/users/mock-login` 으로 발급한다. "
                 "인증이 필요 없는 API는 목록의 인증 칸에 표시"],
        ["데이터 형식", "JSON (UTF-8)"],
        ["시각", "UTC, ISO-8601 (예: `2026-09-22T08:12:44.101Z`). 화면 표시는 Asia/Seoul"],
        ["금액", "원(KRW), JSON number. 소수점이 붙어 올 수 있다 (예: `80000.0000`)"],
    ]))
    w("\n### 공통 헤더\n")
    w(md_table(["헤더", "방향", "설명"], [
        ["`Authorization`", "요청", "인증이 필요한 API에 `Bearer <accessToken>`"],
        ["`Idempotency-Key`", "요청", "주문 접수에만 필수. 같은 키로 같은 내용을 다시 보내면 기존 주문이 돌아온다"],
        ["`X-Trace-Id`", "응답", "모든 응답에 붙는다. 에러 본문의 `traceId`와 같은 값이다"],
    ]))
    w("\n### 에러 응답\n")
    w("성공하지 못한 요청은 모두 같은 형식의 본문을 돌려준다.\n")
    w("```json\n" + pretty(spec.example({"$ref": "#/components/schemas/ErrorResponse"})) + "\n```\n")
    w(md_table(["필드", "타입", "설명"],
               [[code(r["name"]), esc(r["type"]), esc(r["description"])]
                for r in spec.fields({"$ref": "#/components/schemas/ErrorResponse"})]))
    w("\n### 에러 코드\n")
    w(md_table(["코드", "HTTP", "메시지", "발생 API"], [
        [code(e["code"]), esc(e["status"]), esc(e["message"]), esc(error_where(spec, e["code"]))]
        for e in spec.error_codes.values()]))

    w("\n## 2. API 목록\n")
    w(md_table(["No", "분류", "API", "Method", "URI", "인증"], [
        [esc(op["no"]), esc(op["tag"]), f"[{esc(op['summary'])}](#api-{op['no']})", code(op["method"]),
         code(op["path"]), "불필요" if op["public"] else "필요"]
        for op in spec.ops]))

    w("\n## 3. API 상세\n")
    current = None
    for op in spec.ops:
        if op["tag"] != current:
            current = op["tag"]
            w(f"### {current}\n")
            if op["tag_description"]:
                w(op["tag_description"] + "\n")
        w(f'<a id="api-{op["no"]}"></a>\n')
        w(f"#### {op['no']}. {op['summary']}\n")
        w(f"`{op['method']} {op['path']}` · 인증 {'불필요' if op['public'] else '필요'}\n")
        w(op["description"] + "\n")

        req = spec.request_rows(op)
        if req:
            w("**요청**\n")
            w(md_table(["이름", "위치", "타입", "필수", "설명", "예시"], [
                [code(r["name"]), esc(r["in"]), esc(r["type"]), "O" if r["required"] else "",
                 esc(r["description"]), code(r["example"])] for r in req]))
            w("")
        if op["request_schema"]:
            w("```json\n" + pretty(spec.example(op["request_schema"])) + "\n```\n")

        w(f"**응답** — {spec.response_label(op)}\n")
        if op["response_schema"]:
            w(md_table(["필드", "타입", "설명", "예시"], [
                [code(r["name"]), esc(r["type"]), esc(r["description"]), code(r["example"])]
                for r in spec.fields(op["response_schema"])]))
            w("\n```json\n" + pretty(spec.example(op["response_schema"])) + "\n```\n")

        errors = spec.errors_of(op)
        if errors:
            w("**에러**\n")
            w(md_table(["HTTP", "코드", "설명"],
                       [[esc(e["status"]), code(e["code"]), esc(e["message"])] for e in errors]))
            w("")
    return "\n".join(L).rstrip() + "\n"


def error_where(spec, code_name):
    if code_name == "INTERNAL_ERROR":
        return "모든 API (처리하지 못한 예외)"
    nos = [str(op["no"]) for op in spec.ops if code_name in op["errors"]]
    if not nos:
        return "—"
    if len(nos) == len(spec.ops):
        return "모든 API"
    return ", ".join(nos)


# ── Excel ─────────────────────────────────────────────────────────

FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FILL = PatternFill("solid", fgColor="DDE4EE")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def visual_len(text):
    # 한글·전각 문자는 영문 두 칸 정도를 차지한다.
    return sum(1.9 if ord(ch) > 0x2E80 else 1.0 for ch in text)


def lines_needed(text, width):
    usable = max(width - 1.5, 1)
    return sum(max(1, math.ceil(visual_len(part) / usable)) for part in text.split("\n"))


class Sheet:
    """행 단위로 쓰는 얇은 래퍼. 줄바꿈 높이를 직접 계산한다(Excel은 열 때 다시 맞추지 않는다)."""

    def __init__(self, ws, widths):
        self.ws, self.widths, self.row = ws, widths, 1
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    def put(self, values, bold=False, fill=None, merge_from=None, align=None, border=True, size=10):
        """values 를 A열부터 쓴다. merge_from 이 주어지면 그 열부터 끝 열까지 병합한다."""
        r = self.row
        last = len(self.widths)
        height_lines = 1
        for i, v in enumerate(values):
            if merge_from is not None and i > merge_from:
                break
            c = self.ws.cell(row=r, column=i + 1)
            text = cell_text(v)
            if isinstance(v, int) and not isinstance(v, bool):
                c.value = v  # No, HTTP 상태코드는 숫자로 둔다
            else:
                c.value = text
                c.data_type = "s"  # '=' 로 시작하는 설명이 수식으로 읽히지 않게 한다
            c.font = Font(name=FONT, bold=bold, size=size)
            c.alignment = align or WRAP
            span = self.widths[i] if merge_from != i else sum(self.widths[i:])
            height_lines = max(height_lines, lines_needed(text, span))
        if merge_from is not None and merge_from < last - 1:
            self.ws.merge_cells(start_row=r, start_column=merge_from + 1, end_row=r, end_column=last)
        cols = last if merge_from is not None else len(values)
        for col in range(1, cols + 1):
            c = self.ws.cell(row=r, column=col)
            if border:
                c.border = BORDER
            if fill:
                c.fill = fill
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, size=size)
        # Excel 행 높이 상한은 409pt 다. 넘는 예시 JSON은 셀 안에서 스크롤해 본다.
        self.ws.row_dimensions[r].height = min(409, max(15, 13.5 * height_lines + 2))
        self.row += 1
        return r

    def gap(self):
        self.row += 1


def render_xlsx(spec, path):
    wb = Workbook()

    # 개요
    ov = Sheet(wb.active, [18, 100])
    ov.ws.title = "개요"
    ov.put(["WTS 모의투자 API 명세서"], bold=True, border=False, size=14)
    ov.gap()
    for k, v in [
        ("버전", spec.version),
        ("API 수", f"{len(spec.ops)}개"),
        ("Base URL", f"{spec.base_url} (Gateway). 각 서비스를 직접 호출하지 않는다"),
        ("인증", "Authorization: Bearer <accessToken>. 토큰은 POST /api/users/mock-login 으로 발급한다"),
        ("데이터 형식", "JSON (UTF-8)"),
        ("시각", "UTC, ISO-8601 (예: 2026-09-22T08:12:44.101Z). 화면 표시는 Asia/Seoul"),
        ("금액", "원(KRW), JSON number. 소수점이 붙어 올 수 있다 (예: 80000.0000)"),
        ("원본", "컨트롤러(@Operation)와 DTO(@Schema) → docs/api/openapi.yaml → 이 파일. 직접 고치지 않는다"),
        ("다시 만들기", "./infra/scripts/generate-openapi.sh (서비스 기동 필요) "
                       "또는 python3 infra/scripts/render-api-spec.py docs/api/openapi.yaml"),
        ("범위 밖", "WebSocket(/ws/market)과 Kafka 이벤트 계약은 docs/api/README.md 참고"),
    ]:
        ov.put([k, v], bold=False)
        ov.ws.cell(row=ov.row - 1, column=1).font = Font(name=FONT, bold=True, size=10)
        ov.ws.cell(row=ov.row - 1, column=1).fill = HEAD_FILL
    ov.gap()
    ov.put(["공통 헤더"], bold=True, border=False, size=11)
    ov.put(["헤더", "설명"], bold=True, fill=HEAD_FILL)
    for k, v in [("Authorization (요청)", "인증이 필요한 API에 Bearer <accessToken>"),
                 ("Idempotency-Key (요청)", "주문 접수에만 필수. 같은 키로 같은 내용을 다시 보내면 기존 주문이 돌아온다"),
                 ("X-Trace-Id (응답)", "모든 응답에 붙는다. 에러 본문의 traceId와 같은 값이다")]:
        ov.put([k, v])

    # API 목록
    ls = Sheet(wb.create_sheet("API 목록"), [6, 12, 20, 9, 38, 8, 26, 60])
    ls.put(["No", "분류", "API", "Method", "URI", "인증", "성공 응답", "설명"], bold=True, fill=HEAD_FILL)
    list_rows = {}
    for op in spec.ops:
        list_rows[op["no"]] = ls.put([op["no"], op["tag"], op["summary"], op["method"], op["path"],
                                      "불필요" if op["public"] else "필요", spec.response_label(op),
                                      op["description"]])
        for col in (1, 4, 6):
            ls.ws.cell(row=list_rows[op["no"]], column=col).alignment = CENTER
    ls.ws.freeze_panes = "A2"
    ls.ws.auto_filter.ref = f"A1:H{ls.row - 1}"

    # API 상세 — API마다 블록 하나. 요청·응답·에러를 같은 6열 격자에 쓴다.
    dt = Sheet(wb.create_sheet("API 상세"), [32, 9, 22, 6, 62, 30])
    detail_rows = {}
    for op in spec.ops:
        detail_rows[op["no"]] = dt.put([f"{op['no']}. {op['summary']}"], bold=True, fill=TITLE_FILL,
                                       merge_from=0, size=11)
        dt.put(["Method / URI", f"{op['method']}  {op['path']}"], merge_from=1)
        dt.put(["분류 / 인증", f"{op['tag']} · 인증 {'불필요' if op['public'] else '필요'}"], merge_from=1)
        dt.put(["설명", op["description"]], merge_from=1)
        for r in (dt.row - 3, dt.row - 2, dt.row - 1):
            dt.ws.cell(row=r, column=1).font = Font(name=FONT, bold=True, size=10)
            dt.ws.cell(row=r, column=1).fill = HEAD_FILL

        req = spec.request_rows(op)
        if req:
            dt.put(["요청"], bold=True, merge_from=0, border=False)
            dt.put(["이름", "위치", "타입", "필수", "설명", "예시"], bold=True, fill=HEAD_FILL)
            for r in req:
                row = dt.put([r["name"], r["in"], r["type"], "O" if r["required"] else "",
                              r["description"], cell_text(r["example"])])
                dt.ws.cell(row=row, column=4).alignment = CENTER
            if op["request_schema"]:
                dt.put(["요청 예시", pretty(spec.example(op["request_schema"]))], merge_from=1)
                dt.ws.cell(row=dt.row - 1, column=1).font = Font(name=FONT, bold=True, size=10)

        dt.put([f"응답 — {spec.response_label(op)}"], bold=True, merge_from=0, border=False)
        if op["response_schema"]:
            dt.put(["필드", "", "타입", "", "설명", "예시"], bold=True, fill=HEAD_FILL)
            for r in spec.fields(op["response_schema"]):
                dt.put([r["name"], "", r["type"], "", r["description"], cell_text(r["example"])])
            dt.put(["응답 예시", pretty(spec.example(op["response_schema"]))], merge_from=1)
            dt.ws.cell(row=dt.row - 1, column=1).font = Font(name=FONT, bold=True, size=10)

        errors = spec.errors_of(op)
        if errors:
            dt.put(["에러"], bold=True, merge_from=0, border=False)
            dt.put(["코드", "HTTP", "", "", "설명", ""], bold=True, fill=HEAD_FILL)
            for e in errors:
                row = dt.put([e["code"], e["status"], "", "", e["message"], ""])
                dt.ws.cell(row=row, column=2).alignment = CENTER
        dt.gap()
        dt.gap()

    # 목록 → 상세 바로가기
    for no, row in list_rows.items():
        c = ls.ws.cell(row=row, column=3)
        # target("#...") 이 아니라 location 이어야 통합문서 안 이동으로 인식된다.
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'API 상세'!A{detail_rows[no]}")
        c.font = Font(name=FONT, size=10, color="1F4E79", underline="single")

    # 에러 코드
    er = Sheet(wb.create_sheet("에러 코드"), [28, 8, 44, 40])
    er.put(["코드", "HTTP", "메시지", "발생 API (No)"], bold=True, fill=HEAD_FILL)
    for e in spec.error_codes.values():
        row = er.put([e["code"], e["status"], e["message"], error_where(spec, e["code"])])
        er.ws.cell(row=row, column=2).alignment = CENTER
    er.ws.freeze_panes = "A2"

    wb.save(path)


# ── PDF ───────────────────────────────────────────────────────────

# 한글 TrueType 글꼴. reportlab 은 PostScript 윤곽선(OTF/CFF, 예: AppleSDGothicNeo, Noto Sans CJK)을
# 읽지 못한다. 굵은 글꼴이 없으면 보통 글꼴로 대신하고 제목은 크기·배경으로 구분한다.
# 다른 글꼴을 쓰려면 WTS_PDF_FONT(, WTS_PDF_FONT_BOLD) 에 .ttf 경로를 준다.
PDF_FONTS = [
    (os.environ.get("WTS_PDF_FONT"), os.environ.get("WTS_PDF_FONT_BOLD")),
    ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf",            # Linux: fonts-nanum
     "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
    ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),  # Windows: 맑은 고딕
    ("/System/Library/Fonts/Supplemental/AppleGothic.ttf", None),  # macOS
]


def render_pdf(spec, path):
    """PDF 를 만든다. 만들지 못하면 이유를 돌려준다(md/xlsx 생성은 막지 않는다)."""
    try:
        from xml.sax.saxutils import escape
        from reportlab import rl_config
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (CondPageBreak, PageBreak, Paragraph, SimpleDocTemplate,
                                        Spacer, Table, TableStyle, XPreformatted)
    except ImportError:
        return "reportlab 이 없다 (pip install reportlab)"

    regular, bold = next(((r, b) for r, b in PDF_FONTS if r and os.path.exists(r)), (None, None))
    if not regular:
        return "한글 TrueType 글꼴을 찾지 못했다 (WTS_PDF_FONT 에 .ttf 경로를 준다)"
    pdfmetrics.registerFont(TTFont("KR", regular))
    pdfmetrics.registerFont(TTFont("KR-B", bold if bold and os.path.exists(bold) else regular))
    pdfmetrics.registerFontFamily("KR", normal="KR", bold="KR-B", italic="KR", boldItalic="KR-B")
    rl_config.invariant = 1  # 생성 시각·ID를 고정한다. 명세가 같으면 PDF 바이트도 같다

    INK, MUTED = colors.HexColor("#1F2328"), colors.HexColor("#57606A")
    LINE, HEAD_BG = colors.HexColor("#D0D7DE"), colors.HexColor("#F2F4F7")
    BAND, CODE_BG = colors.HexColor("#DDE4EE"), colors.HexColor("#F6F8FA")
    LINK = colors.HexColor("#1F4E79")

    def style(name, **kw):
        base = dict(fontName="KR", fontSize=8.5, leading=11.5, textColor=INK)
        base.update(kw)
        return ParagraphStyle(name, **base)

    S = {
        "title": style("title", fontName="KR-B", fontSize=20, leading=26, spaceAfter=4),
        "sub": style("sub", fontSize=9.5, leading=13, textColor=MUTED, spaceAfter=10),
        "h1": style("h1", fontName="KR-B", fontSize=15, leading=20, spaceBefore=6, spaceAfter=8),
        "h2": style("h2", fontName="KR-B", fontSize=11.5, leading=15, spaceBefore=10, spaceAfter=5),
        "api": style("api", fontName="KR-B", fontSize=11, leading=14),
        "label": style("label", fontName="KR-B", fontSize=9, leading=12, spaceBefore=7, spaceAfter=3),
        "meta": style("meta", fontSize=9, leading=12, textColor=MUTED, spaceBefore=3, spaceAfter=4),
        "body": style("body", fontSize=9, leading=13, spaceAfter=4),
        "note": style("note", fontSize=8.5, leading=12, textColor=MUTED, backColor=CODE_BG,
                      borderPadding=6, leftIndent=6, rightIndent=6, spaceBefore=4, spaceAfter=12),
        "cell": style("cell"),
        # 필드명·타입·예시·URI 는 공백 없이 길다. CJK 줄바꿈은 아무 글자 사이에서나 끊을 수 있다.
        "code": style("code", wordWrap="CJK"),
        "head": style("head", fontName="KR-B", textColor=INK),
        "pre": style("pre", fontSize=7.5, leading=9.8, backColor=CODE_BG, borderPadding=6,
                     leftIndent=6, rightIndent=6, spaceBefore=6, spaceAfter=8),
    }

    def P(text, st="cell"):
        return Paragraph(escape(cell_text(text)).replace("\n", "<br/>"), S[st])

    def table(headers, rows, widths, code_cols=(), center_cols=()):
        data = [[P(h, "head") for h in headers]]
        for r in rows:
            data.append([c if not isinstance(c, (str, int, float)) else
                         P(c, "code" if i in code_cols else "cell") for i, c in enumerate(r)])
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        cmds = [("GRID", (0, 0), (-1, -1), 0.5, LINE),
                ("BACKGROUND", (0, 0), (-1, 0), HEAD_BG),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        t.setStyle(TableStyle(cmds))
        return t

    def heading(text, st, key, level):
        h = Paragraph(f'<a name="{key}"/>{escape(text)}', S[st])
        h.outline = (level, text, key)  # afterFlowable 이 PDF 책갈피로 만든다
        return h

    def pre(obj):
        return XPreformatted(escape(pretty(obj)), S["pre"])

    W = A4[0] - 30 * mm  # 본문 폭 (좌우 여백 15mm)
    story = []
    add = story.append

    add(Paragraph("WTS 모의투자 API 명세서", S["title"]))
    add(Paragraph(escape(f"버전 {spec.version} · API {len(spec.ops)}개 · Base URL {spec.base_url}"), S["sub"]))
    add(Paragraph(
        "이 문서는 생성물이다. 원본은 컨트롤러(@Operation)와 DTO(@Schema)이고 docs/api/openapi.yaml 을 거쳐 "
        "./infra/scripts/generate-openapi.sh 가 만든다. 같은 내용이 api-spec.md, api-spec.xlsx 에 있다. "
        "실시간 시세 WebSocket(/ws/market)과 Kafka 이벤트 계약은 docs/api/README.md 를 본다.", S["note"]))

    add(heading("1. 공통", "h1", "sec-common", 0))
    add(table(["항목", "내용"], [
        ["Base URL", f"{spec.base_url} (Gateway). 각 서비스를 직접 호출하지 않는다"],
        ["인증", "Authorization: Bearer {accessToken}. 토큰은 POST /api/users/mock-login 으로 발급한다. "
                 "인증이 필요 없는 API는 API 목록의 인증 칸에 표시"],
        ["데이터 형식", "JSON (UTF-8)"],
        ["시각", "UTC, ISO-8601 (예: 2026-09-22T08:12:44.101Z). 화면 표시는 Asia/Seoul"],
        ["금액", "원(KRW), JSON number. 소수점이 붙어 올 수 있다 (예: 80000.0000)"],
    ], [28 * mm, W - 28 * mm]))

    add(heading("공통 헤더", "h2", "sec-headers", 1))
    add(table(["헤더", "방향", "설명"], [
        ["Authorization", "요청", "인증이 필요한 API에 Bearer {accessToken}"],
        ["Idempotency-Key", "요청", "주문 접수에만 필수. 같은 키로 같은 내용을 다시 보내면 기존 주문이 돌아온다"],
        ["X-Trace-Id", "응답", "모든 응답에 붙는다. 에러 본문의 traceId와 같은 값이다"],
    ], [34 * mm, 14 * mm, W - 48 * mm], code_cols=(0,)))

    add(heading("에러 응답", "h2", "sec-error-body", 1))
    add(Paragraph("성공하지 못한 요청은 모두 같은 형식의 본문을 돌려준다.", S["body"]))
    error_ref = {"$ref": "#/components/schemas/ErrorResponse"}
    add(pre(spec.example(error_ref)))
    add(table(["필드", "타입", "설명"],
              [[r["name"], r["type"], r["description"]] for r in spec.fields(error_ref)],
              [30 * mm, 32 * mm, W - 62 * mm], code_cols=(0,)))

    add(heading("에러 코드", "h2", "sec-error-codes", 1))
    add(table(["코드", "HTTP", "메시지", "발생 API (No)"],
              [[e["code"], e["status"], e["message"], error_where(spec, e["code"])]
               for e in spec.error_codes.values()],
              [52 * mm, 13 * mm, W - 107 * mm, 42 * mm], code_cols=(0, 3)))

    add(PageBreak())
    add(heading("2. API 목록", "h1", "sec-list", 0))
    link = style("link", textColor=LINK)
    rows = [[op["no"], op["tag"],
             Paragraph(f'<a href="#api-{op["no"]}">{escape(op["summary"])}</a>', link),
             op["method"], op["path"], "불필요" if op["public"] else "필요"] for op in spec.ops]
    add(table(["No", "분류", "API", "Method", "URI", "인증"], rows,
              [9 * mm, 19 * mm, 34 * mm, 16 * mm, W - 92 * mm, 14 * mm], code_cols=(3, 4)))

    add(PageBreak())
    add(heading("3. API 상세", "h1", "sec-detail", 0))
    current = None
    for op in spec.ops:
        if op["tag"] != current:
            current = op["tag"]
            add(CondPageBreak(60 * mm))
            add(heading(current, "h2", f"tag-{current}", 1))
            if op["tag_description"]:
                add(Paragraph(escape(op["tag_description"]), S["meta"]))
        add(CondPageBreak(45 * mm))  # 제목만 페이지 끝에 홀로 남지 않게 한다
        band = Table([[heading(f"{op['no']}. {op['summary']}", "api", f"api-{op['no']}", 2)]],
                     colWidths=[W], hAlign="LEFT")
        band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BAND),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                  ("TOPPADDING", (0, 0), (-1, -1), 4),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        band.outline = band._cellvalues[0][0].outline
        add(band)
        add(Paragraph(escape(f"{op['method']}  {op['path']}  ·  {op['tag']}  ·  "
                             f"인증 {'불필요' if op['public'] else '필요'}"), S["meta"]))
        add(Paragraph(escape(op["description"]), S["body"]))

        req = spec.request_rows(op)
        if req:
            add(Paragraph("요청", S["label"]))
            add(table(["이름", "위치", "타입", "필수", "설명", "예시"],
                      [[r["name"], r["in"], r["type"], "O" if r["required"] else "",
                        r["description"], cell_text(r["example"])] for r in req],
                      [32 * mm, 13 * mm, 30 * mm, 9 * mm, W - 122 * mm, 38 * mm],
                      code_cols=(0, 5)))
            if op["request_schema"]:
                add(pre(spec.example(op["request_schema"])))

        add(Paragraph(escape("응답 — " + spec.response_label(op)), S["label"]))
        if op["response_schema"]:
            add(table(["필드", "타입", "설명", "예시"],
                      [[r["name"], r["type"], r["description"], cell_text(r["example"])]
                       for r in spec.fields(op["response_schema"])],
                      [52 * mm, 30 * mm, W - 120 * mm, 38 * mm], code_cols=(0, 3)))
            add(pre(spec.example(op["response_schema"])))

        errors = spec.errors_of(op)
        if errors:
            add(Paragraph("에러", S["label"]))
            add(table(["HTTP", "코드", "설명"],
                      [[e["status"], e["code"], e["message"]] for e in errors],
                      [14 * mm, 56 * mm, W - 70 * mm], code_cols=(1,)))
        add(Spacer(1, 8 * mm))

    class Doc(SimpleDocTemplate):
        def afterFlowable(self, flowable):
            outline = getattr(flowable, "outline", None)
            if outline:
                level, title, key = outline
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(title, key, level=level, closed=level >= 1)

    def footer(canv, doc):
        canv.saveState()
        canv.setStrokeColor(LINE)
        canv.setLineWidth(0.5)
        canv.line(15 * mm, 12 * mm, A4[0] - 15 * mm, 12 * mm)
        canv.setFont("KR", 7.5)
        canv.setFillColor(MUTED)
        canv.drawString(15 * mm, 8 * mm, f"WTS 모의투자 API 명세서 · {spec.version}")
        canv.drawRightString(A4[0] - 15 * mm, 8 * mm, str(doc.page))
        canv.restoreState()

    doc = Doc(path, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
              topMargin=15 * mm, bottomMargin=18 * mm,
              title="WTS 모의투자 API 명세서", author="SeSac Team3", subject=f"API {spec.version}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return None


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    with open(src, encoding="utf-8") as f:
        spec = Spec(yaml.safe_load(f))
    out_dir = os.path.dirname(os.path.abspath(src))
    md_path = os.path.join(out_dir, "api-spec.md")
    xlsx_path = os.path.join(out_dir, "api-spec.xlsx")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(render_md(spec))
    render_xlsx(spec, xlsx_path)
    print(md_path)
    print(xlsx_path)
    pdf_path = os.path.join(out_dir, "api-spec.pdf")
    skipped = render_pdf(spec, pdf_path)
    print(f"PDF 건너뜀: {skipped}" if skipped else pdf_path)
    print(f"  API {len(spec.ops)}개 · 에러 코드 {len(spec.error_codes)}개")


if __name__ == "__main__":
    main()
